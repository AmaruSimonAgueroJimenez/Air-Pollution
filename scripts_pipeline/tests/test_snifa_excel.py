"""Offline checks of SNIFA Excel identity, source semantics and known samples."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
import hashlib
import importlib.util
import os
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from _snifa_excel import parse_excel


STATIONS = [
    {"id": "el_penon", "aliases": ["El Peñón", "Central El Peñón"]},
    {"id": "sapu", "aliases": ["SAPU"]},
    {"id": "quinel", "aliases": ["Quinel"]},
]
SAMPLE_ROOT = Path(os.environ.get("SNIFA_EXCEL_SAMPLE_ROOT", str(
    Path.home() / ".codex/.chatgpt-projects/g-p-6ab00989aedc8191997dcf9bcc53cbdf/investigacion_sinca_20260922")))
NETWORK_SAMPLE_ROOT = Path(os.environ.get("SNIFA_NETWORK_SAMPLE_ROOT",
    "/Volumes/Datos/Asesorias_Data/AirPollution/data/snifa_adicional"))


class ExcelTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="snifa-excel-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def workbook(self, filename="SAPU.xlsx", sheets=None, epoch=None):
        """Write small disposable fixtures only; production parsers never author."""
        workbook = Workbook()
        workbook.remove(workbook.active)
        if epoch is not None:
            workbook.epoch = epoch
        for spec in sheets or [{}]:
            sheet = workbook.create_sheet(spec.get("name", "NO2"))
            sheet.append(["ESTACIÓN", spec.get("station", "SAPU")])
            sheet.append(["UNIDAD", spec.get("unit", "ppb")])
            sheet.append(["PERÍODO", spec.get("period", "01 al 30 de junio del 2025")])
            sheet.append(["Fecha", *spec.get("hours", range(1, 25)), "MEDIA"])
            for day, values in spec.get("data", [(datetime(2025, 6, 1), list(range(24)))]):
                sheet.append([day, *values, 99999])
            sheet.append(["MEDIA", *([123] * 24)])
        path = self.root / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()
        return path

    def test_codes_blanks_zero_decimal_and_locator_are_preserved(self):
        values = [0, None, "2.e", "1,25", "<0.1", -2, "=1+2", False] + [3] * 16
        path = self.workbook(sheets=[{"data": [(20250601, values)]}])
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = parse_excel(path, STATIONS)
        self.assertEqual(len(result["rows"]), 24)
        self.assertEqual(result["issues"], [])
        first = result["rows"][0]
        self.assertEqual(first, {"station_id": "sapu", "pollutant": "no2", "date": "2025-06-01",
                                 "hour_label": "1", "value": 0.0, "raw_value": "0", "unit": "ppb",
                                 "quality_code": "", "source_locator": "NO2!B5", "resolution": "hourly"})
        self.assertIsNone(result["rows"][1]["value"])
        self.assertEqual(result["rows"][1]["raw_value"], "")
        self.assertEqual(result["rows"][2]["quality_code"], "2.e")
        self.assertEqual(result["rows"][3]["value"], 1.25)
        self.assertEqual(result["rows"][3]["raw_value"], "1,25")
        self.assertEqual(result["rows"][4]["quality_code"], "<0.1")
        self.assertEqual(result["rows"][5]["value"], -2)
        self.assertEqual(result["rows"][6]["quality_code"], "FORMULA_NO_CACHED_VALUE")
        self.assertEqual(result["rows"][6]["raw_value"], "=1+2")
        self.assertIsNone(result["rows"][7]["value"])
        self.assertEqual(result["rows"][-1]["hour_label"], "24")
        self.assertEqual(result["rows"][-1]["date"], "2025-06-01")
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_rolling_particulate_no_and_nox_never_become_hourly_no2(self):
        sheets = [{"name": name} for name in ("NO2", "O3 8 Hrs", "CO 8 horas", "NO", "NOx", "MP10", "PM2.5")]
        result = parse_excel(self.workbook(sheets=sheets), STATIONS)
        self.assertEqual(len(result["rows"]), 24)
        self.assertEqual({row["pollutant"] for row in result["rows"]}, {"no2"})

    def test_all_four_gases_and_unicode_subscripts(self):
        result = parse_excel(self.workbook(sheets=[{"name": name} for name in ("CO", "SO₂", "NO₂", "O₃")]), STATIONS)
        self.assertEqual(Counter(row["pollutant"] for row in result["rows"]),
                         Counter({"co": 24, "so2": 24, "no2": 24, "o3": 24}))

    def test_station_identity_missing_conflicting_and_multiple_station_workbook(self):
        path = self.workbook("anonymous.xlsx", [{"station": None}])
        result = parse_excel(path, [STATIONS[1]])
        self.assertEqual(result["rows"], [])
        self.assertTrue(any("estación" in issue for issue in result["issues"]))
        path = self.workbook("Quinel.xlsx", [{"station": "SAPU"}])
        self.assertEqual(parse_excel(path, STATIONS)["rows"], [])
        path = self.workbook("SAPU.xlsx", [{"station": "Otra estación ajena al inventario"}])
        self.assertEqual(parse_excel(path, STATIONS)["rows"], [])
        path = self.workbook("SAPU_Quinel.xlsx", [{"name": "SAPU NO2", "station": "SAPU"},
                                                 {"name": "Quinel NO2", "station": "Quinel"}])
        self.assertEqual(Counter(row["station_id"] for row in parse_excel(path, STATIONS)["rows"]),
                         Counter({"sapu": 24, "quinel": 24}))

    def test_parent_station_directory_and_accent_normalization(self):
        path = self.workbook("el_penon/anonymous.xlsx", [{"station": None}])
        result = parse_excel(path, STATIONS)
        self.assertEqual({row["station_id"] for row in result["rows"]}, {"el_penon"})
        path = self.workbook("anonymous2.xlsx", [{"station": "Central EL PENÓN"}])
        self.assertEqual({row["station_id"] for row in parse_excel(path, STATIONS)["rows"]}, {"el_penon"})

    def test_unknown_unit_never_defaults_to_ppb(self):
        path = self.workbook(sheets=[{"unit": None}])
        result = parse_excel(path, STATIONS)
        self.assertEqual({row["unit"] for row in result["rows"]}, {"unknown"})
        self.assertTrue(any("unidad" in issue for issue in result["issues"]))
        path = self.workbook(sheets=[{"unit": "µg/m3N"}])
        self.assertEqual({row["unit"] for row in parse_excel(path, STATIONS)["rows"]}, {"µg/m3N"})

    def test_serial_dates_1900_and_1904_and_hhmm_labels(self):
        hours = [f"{hour:02}00" for hour in range(1, 25)]
        path = self.workbook(sheets=[{"hours": hours, "data": [(to_excel(datetime(2025, 6, 1)), [1] * 24)]}])
        result = parse_excel(path, STATIONS)
        self.assertEqual(result["rows"][0]["date"], "2025-06-01")
        self.assertEqual(result["rows"][0]["hour_label"], "0100")
        self.assertEqual(result["rows"][-1]["hour_label"], "2400")
        path = self.workbook(epoch=CALENDAR_MAC_1904, sheets=[{"data": [(
            to_excel(datetime(2025, 6, 1), CALENDAR_MAC_1904), [1] * 24)]}])
        self.assertEqual(parse_excel(path, STATIONS)["rows"][0]["date"], "2025-06-01")

    def test_declared_period_excludes_template_extra_day_but_keeps_legitimate_zero(self):
        data = [(datetime(2025, 6, 29), [0] * 24), (datetime(2025, 6, 30), [0] * 24),
                (datetime(2025, 7, 1), [0] * 24)]
        path = self.workbook(sheets=[{"data": data}])
        result = parse_excel(path, STATIONS)
        self.assertEqual(len(result["rows"]), 48)
        self.assertTrue(all(row["value"] == 0 for row in result["rows"]))
        self.assertTrue(any("período declarado" in issue for issue in result["issues"]))
        result = parse_excel(path, STATIONS, "2025-06-30", "2025-06-30")
        self.assertEqual(len(result["rows"]), 24)
        self.assertEqual({row["date"] for row in result["rows"]}, {"2025-06-30"})

    def test_zero_based_hours_and_duplicate_dates_remain_native(self):
        path = self.workbook(sheets=[{"hours": range(24), "data": [(20250601, [1] * 24), (20250601, [2] * 24)]}])
        rows = parse_excel(path, STATIONS)["rows"]
        self.assertEqual(len(rows), 48)
        self.assertEqual(rows[0]["hour_label"], "0")
        self.assertEqual(rows[23]["hour_label"], "23")
        self.assertNotEqual(rows[0]["source_locator"], rows[24]["source_locator"])

    def test_bad_file_is_reported_and_inverted_bounds_rejected(self):
        path = self.root / "SAPU.xlsx"
        path.write_text("not an Excel workbook")
        self.assertEqual(parse_excel(path, STATIONS)["rows"], [])
        self.assertTrue(parse_excel(path, STATIONS)["issues"])
        with self.assertRaises(ValueError):
            parse_excel(path, STATIONS, "2026-01-02", "2026-01-01")

    def test_xlsx_bytes_under_xls_name_keep_original_values_and_formula_flags(self):
        original = self.workbook("SAPU.xlsx", [{"data": [(20250601, [1, "2.e", "=1+2"] + [0] * 21)]}])
        expected = parse_excel(original, STATIONS)
        mislabeled = self.root / "SAPU.XLS"
        mislabeled.write_bytes(original.read_bytes())
        before = hashlib.sha256(mislabeled.read_bytes()).hexdigest()
        self.assertTrue(mislabeled.read_bytes().startswith(b"PK"))
        result = parse_excel(mislabeled, STATIONS)
        self.assertEqual(result, expected)
        self.assertEqual(len(result["rows"]), 24)
        self.assertEqual(result["rows"][2]["quality_code"], "FORMULA_NO_CACHED_VALUE")
        self.assertEqual(hashlib.sha256(mislabeled.read_bytes()).hexdigest(), before)
        self.assertTrue(mislabeled.exists())


class ColumnarTest(unittest.TestCase):
    STATIONS = [{"id": "charrua", "aliases": ["Charrúa", "Charrúa Sur"]},
                {"id": "progreso", "aliases": ["Progreso"]}]

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="snifa-columns-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def workbook(self, specifications=None, filename="Datos Columna.xlsx"):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for spec in specifications or [{}]:
            station = spec.get("station", "CHARRUA")
            sheet = workbook.create_sheet(spec.get("sheet", "DATOS " + station))
            sheet.append(["MONITOREOS ESTACIÓN " + station])
            sheet.append(["Coordenadas UTM WGS84"])
            sheet.append([])
            sheet.append([spec.get("period", "PERÍODO ABRIL DE 2020")])
            sheet.append([])
            kinds = ["Concentración Promedios horarios"] * 8
            kinds[2] = kinds[4] = "Concentración Móvil 8 horas"
            if spec.get("events"):
                kinds = ["Código de invalidación"] * 8
            sheet.append(["Fecha y hora", *spec.get("kinds", kinds)])
            labels = ["SO2\nppb", "O3\nppb", "O3\nppb", "CO\nppm", "CO móvil\nppm",
                      "NO\nppb", "NO2\nppb", "NOX\nppb"]
            if spec.get("events"):
                labels = ["SO2", "O3", "O3 móvil", "CO", "CO móvil", "NO", "NO2", "NOX"]
            sheet.append([None, *spec.get("labels", labels)])
            for stamp, values in spec.get("data", [(202004010000, [1, 2, 999, 0.1, 888, 444, 3, 555])]):
                sheet.append([stamp, *values])
        path = self.root / filename
        workbook.save(path)
        workbook.close()
        return path

    def test_columns_use_type_and_parameter_not_position_or_nox(self):
        path = self.workbook()
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = parse_excel(path, self.STATIONS)
        self.assertEqual(result["issues"], [])
        self.assertEqual({row["pollutant"]: row["value"] for row in result["rows"]},
                         {"so2": 1, "o3": 2, "co": 0.1, "no2": 3})
        self.assertEqual({row["pollutant"]: row["unit"] for row in result["rows"]},
                         {"so2": "ppb", "o3": "ppb", "co": "ppm", "no2": "ppb"})
        self.assertEqual({row["hour_label"] for row in result["rows"]}, {"0000"})
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_native_compact_datetime_serial_and_2400_timestamps(self):
        stamps = [202004010000, datetime(2020, 4, 1, 1), to_excel(datetime(2020, 4, 1, 2)),
                  "2020-04-01 03:00:00", "01/04/2020 04:00", 202004012400]
        path = self.workbook([{"data": [(stamp, [1] * 8) for stamp in stamps]}])
        rows = parse_excel(path, self.STATIONS)["rows"]
        self.assertEqual(len(rows), 24)
        self.assertEqual({row["date"] for row in rows}, {"2020-04-01"})
        self.assertEqual([row["hour_label"] for row in rows if row["pollutant"] == "no2"],
                         ["0000", "0100", "0200", "0300", "0400", "2400"])

    def test_month_and_explicit_bounds_reject_extra_dates_and_non_hourly_rows(self):
        data = [(202004010000, [1] * 8), (202004020000, [2] * 8),
                (202005010000, [99] * 8), (202004010030, [88] * 8)]
        result = parse_excel(self.workbook([{"data": data}]), self.STATIONS, "2020-04-02", "2020-04-30")
        self.assertEqual(len(result["rows"]), 4)
        self.assertEqual({row["date"] for row in result["rows"]}, {"2020-04-02"})
        self.assertTrue(any("período declarado" in issue for issue in result["issues"]))
        self.assertTrue(any("sello horario entero" in issue for issue in result["issues"]))

    def test_several_stations_and_unknown_identity_never_use_filename_fallback(self):
        path = self.workbook([{"station": "CHARRUA"}, {"station": "PROGRESO"},
                              {"station": "PUENTES NEGROS"}])
        result = parse_excel(path, self.STATIONS)
        self.assertEqual(Counter(row["station_id"] for row in result["rows"]),
                         Counter({"charrua": 4, "progreso": 4}))
        path = self.workbook([{"station": "PUENTES NEGROS", "sheet": "DATOS CHARRUA"}], "Charrua.xlsx")
        self.assertEqual(parse_excel(path, self.STATIONS)["rows"], [])

    def test_event_codes_join_by_station_timestamp_and_gas_with_cell_provenance(self):
        specs = [
            {"station": "CHARRUA", "data": [(202004010000, [1, 2, 9, 0.1, 8, 4, "2.e", 5])]},
            {"station": "CHARRUA", "sheet": "EVENTOS CHARRUA", "events": True,
             "data": [(datetime(2020, 4, 1), ["2.a", None, None, None, None, None, "3.b", None])]},
            {"station": "PROGRESO", "sheet": "EVENTOS PROGRESO", "events": True,
             "data": [(202004010000, ["WRONG-STATION"] * 8)]},
        ]
        result = parse_excel(self.workbook(specs), self.STATIONS)
        self.assertEqual(len(result["rows"]), 4)
        by = {row["pollutant"]: row for row in result["rows"]}
        self.assertEqual(by["so2"]["value"], 1)
        self.assertEqual(by["so2"]["quality_code"], "2.a")
        self.assertEqual(by["so2"]["source_locator"], "DATOS CHARRUA!B8; EVENTOS CHARRUA!B8")
        self.assertIsNone(by["no2"]["value"])
        self.assertEqual(by["no2"]["quality_code"], "2.e | 3.b")
        self.assertEqual(by["no2"]["raw_value"], "2.e")
        self.assertEqual(by["co"]["quality_code"], "")

    def test_column_units_remain_original_and_ambiguous_parameters_are_excluded(self):
        labels = ["SO2\nµg/m3N", "O3", "O3", "CO\nmg/m3N", "CO móvil", "NO", "NO2", "NOX"]
        result = parse_excel(self.workbook([{"labels": labels}]), self.STATIONS)
        by = {row["pollutant"]: row for row in result["rows"]}
        self.assertEqual(by["so2"]["unit"], "µg/m3N")
        self.assertEqual(by["co"]["unit"], "mg/m3N")
        self.assertEqual(by["no2"]["unit"], "unknown")
        self.assertTrue(any("unidad" in issue for issue in result["issues"]))
        labels[5] = "NO2\nppb"  # Two columns claiming hourly NO2: preserve no guess.
        result = parse_excel(self.workbook([{"labels": labels}]), self.STATIONS)
        self.assertNotIn("no2", {row["pollutant"] for row in result["rows"]})
        self.assertTrue(any("duplicadas/ambiguas" in issue for issue in result["issues"]))


class ValidatedColumnsTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="snifa-validated-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def workbook(self, station="El Peñón", filename="observaciones.xlsx", unit="ppb"):
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Hoja1"
        if station is not None:
            summary.append(["Nombre de la estacion", station])
        sheet = workbook.create_sheet("NO2 data horaria validada")
        sheet.append(["aaaammdd hhmm", unit, "Código de Invalidación de datos (D.S. N°61/2008)"])
        for row in [["20250501 0100", 3.8, ""], ["20250501 0200", 4.3, "2.e"],
                    ["20250501 0300", "ND", "2.b"], ["20250501 2400", 0, ""],
                    ["20250501 0005", 99, ""], ["20250601 0100", 88, ""]]:
            sheet.append(row)
        for name in ("NO2 data cruda", "NOX data horaria validada", "NO data horaria validada"):
            raw = workbook.create_sheet(name)
            raw.append(["aaaammdd hhmm", unit, "Código de Invalidación de datos"])
            raw.append(["20250501 0100", 12345, ""])
        path = self.root / filename
        workbook.save(path)
        workbook.close()
        return path

    def test_hourly_values_codes_zero_and_published_2400_with_no_raw_nox_leakage(self):
        path = self.workbook()
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = parse_excel(path, STATIONS, "2025-05-01", "2025-05-31")
        self.assertEqual(len(result["rows"]), 4)
        self.assertEqual({r["pollutant"] for r in result["rows"]}, {"no2"})
        self.assertEqual({r["station_id"] for r in result["rows"]}, {"el_penon"})
        self.assertEqual([r["hour_label"] for r in result["rows"]], ["0100", "0200", "0300", "2400"])
        self.assertEqual([r["value"] for r in result["rows"]], [3.8, 4.3, None, 0.0])
        self.assertEqual([r["quality_code"] for r in result["rows"]], ["", "2.e", "ND | 2.b", ""])
        self.assertEqual(result["rows"][1]["source_locator"],
                         "NO2 data horaria validada!B3; NO2 data horaria validada!C3")
        self.assertTrue(any("sello horario entero" in issue for issue in result["issues"]))
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_station_evidence_required_and_explicit_unknown_overrides_filename(self):
        for station, filename in [(None, "observaciones.xlsx"), ("Otra estación", "El Peñón.xlsx")]:
            with self.subTest(station=station):
                result = parse_excel(self.workbook(station=station, filename=filename), STATIONS)
                self.assertEqual(result["rows"], [])
                self.assertTrue(any("estación" in issue for issue in result["issues"]))

    def test_unknown_unit_stays_unknown(self):
        result = parse_excel(self.workbook(unit="concentración"), STATIONS)
        self.assertEqual({r["unit"] for r in result["rows"]}, {"unknown"})
        self.assertTrue(any("unidad" in issue for issue in result["issues"]))


class ExistingSampleTest(unittest.TestCase):
    """Optional integration fixtures remain at their existing read-only locations."""

    @unittest.skipUnless((NETWORK_SAMPLE_ROOT / "raw/el_penon/1072705/1286427_Tabla Hr Gases Estacion El Peñon 05-25.xls").exists()
                         and importlib.util.find_spec("python_calamine"),
                         "downloaded XLS with xlrd SST assertion/calamine unavailable")
    def test_real_1286427_xls_sst_fallback_preserves_native_cells_and_codes(self):
        import xlrd

        path = NETWORK_SAMPLE_ROOT / "raw/el_penon/1072705/1286427_Tabla Hr Gases Estacion El Peñon 05-25.xls"
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        with self.assertRaises(AssertionError):
            xlrd.open_workbook(str(path), on_demand=True)
        result = parse_excel(path, STATIONS, "2025-05-01", "2025-05-31")
        self.assertEqual(len(result["rows"]), 2976)
        self.assertTrue(any("Lectura alternativa python-calamine" in issue for issue in result["issues"]))
        self.assertFalse(any("No se pudo leer" in issue for issue in result["issues"]))
        # Cached native precision is retained, not replaced with rounded report text.
        expected = {"co": (0.067852762275, "ppm", 3),
                    "so2": (0.400166666667, "ppb", 6),
                    "no2": (3.8425, "ppb", 5), "o3": (9.333833333333, "ppb", 5)}
        for pollutant, (first, unit, codes) in expected.items():
            rows = [row for row in result["rows"] if row["pollutant"] == pollutant]
            self.assertEqual(len(rows), 744)
            self.assertEqual({row["station_id"] for row in rows}, {"el_penon"})
            self.assertEqual({row["unit"] for row in rows}, {unit})
            self.assertEqual((rows[0]["date"], rows[0]["hour_label"]), ("2025-05-01", "100"))
            self.assertEqual(rows[0]["source_locator"], f"{pollutant.upper()}!C10")
            self.assertAlmostEqual(rows[0]["value"], first)
            invalid = [row for row in rows if row["quality_code"]]
            self.assertEqual(len(invalid), codes)
            self.assertTrue(all(row["value"] is None and row["raw_value"] == "2.e"
                                and row["quality_code"] == "2.e" for row in invalid))
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    @unittest.skipUnless((NETWORK_SAMPLE_ROOT / "raw/el_penon/1072705/1286431_Datos Validados Estación El Peñon 05-25.xls").exists()
                         and importlib.util.find_spec("python_calamine"),
                         "downloaded validated XLS/calamine unavailable")
    def test_real_validated_columns_match_sibling_dates_codes_and_published_rounding(self):
        directory = NETWORK_SAMPLE_ROOT / "raw/el_penon/1072705"
        path = directory / "1286431_Datos Validados Estación El Peñon 05-25.xls"
        matrix = directory / "1286427_Tabla Hr Gases Estacion El Peñon 05-25.xls"
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        validated = parse_excel(path, STATIONS, "2025-05-01", "2025-05-31")
        original = parse_excel(matrix, STATIONS, "2025-05-01", "2025-05-31")
        self.assertEqual(len(validated["rows"]), 2976)
        key = lambda row: (row["station_id"], row["pollutant"], row["date"], int(row["hour_label"]))
        by_key = {key(row): row for row in original["rows"]}
        self.assertEqual(set(by_key), {key(row) for row in validated["rows"]})
        for row in validated["rows"]:
            native = by_key[key(row)]
            self.assertEqual(row["unit"], native["unit"])
            self.assertEqual(row["quality_code"], native["quality_code"])
            if row["value"] is not None and native["value"] is not None:
                tolerance = 0.50001 if row["pollutant"] == "o3" else 0.050001
                self.assertLessEqual(abs(row["value"] - native["value"]), tolerance)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    @unittest.skipUnless((NETWORK_SAMPLE_ROOT / "raw/los_guindos/97026/244839_METEOR_3.XLS").exists(),
                         "downloaded mislabeled XLS sample unavailable")
    def test_real_244839_xlsx_under_xls_name_is_read_without_inventing_gases(self):
        path = NETWORK_SAMPLE_ROOT / "raw/los_guindos/97026/244839_METEOR_3.XLS"
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        with path.open("rb") as stream:
            self.assertEqual(stream.read(2), b"PK")
        result = parse_excel(path, ColumnarTest.STATIONS, "2019-07-01", "2019-09-30")
        self.assertEqual(result["rows"], [])  # The real workbook contains meteorology only.
        self.assertTrue(any("dentro del alcance" in issue for issue in result["issues"]))
        self.assertFalse(any("No se pudo leer" in issue for issue in result["issues"]))
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    @unittest.skipUnless((SAMPLE_ROOT / "el_penon/gases_2026_07.xls").exists()
                         and importlib.util.find_spec("xlrd"), "local XLS sample/xlrd unavailable")
    def test_el_penon_original_counts_units_and_codes(self):
        path = SAMPLE_ROOT / "el_penon/gases_2026_07.xls"
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = parse_excel(path, STATIONS)
        self.assertEqual(result["issues"], [])
        self.assertEqual(len(result["rows"]), 2976)
        expected = {"co": (736, "ppm", {"2.e": 5, "3.b": 3}),
                    "so2": (732, "ppb", {"2.e": 5, "3.b": 3, "2.a": 4}),
                    "no2": (735, "ppb", {"2.e": 5, "3.b": 3, "2.a": 1}),
                    "o3": (735, "ppb", {"2.e": 4, "3.b": 3, "2.a": 2})}
        for pollutant, (count, unit, codes) in expected.items():
            rows = [row for row in result["rows"] if row["pollutant"] == pollutant]
            self.assertEqual(len(rows), 744)
            self.assertEqual(sum(row["value"] is not None for row in rows), count)
            self.assertEqual({row["unit"] for row in rows}, {unit})
            self.assertEqual(Counter(row["quality_code"] for row in rows if row["quality_code"]), Counter(codes))
            self.assertEqual({row["station_id"] for row in rows}, {"el_penon"})
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    @unittest.skipUnless((SAMPLE_ROOT / "otras_estaciones").exists(), "local XLSX samples unavailable")
    def test_sapu_quinel_declared_month_unit_and_no_pm_leakage(self):
        samples = sorted((SAMPLE_ROOT / "otras_estaciones").glob("*.xlsx"))
        if not samples:
            self.skipTest("local XLSX samples unavailable")
        for path in samples:
            with self.subTest(path=path.name):
                before = hashlib.sha256(path.read_bytes()).hexdigest()
                result = parse_excel(path, STATIONS)
                station = "sapu" if "SAPU" in path.name else "quinel"
                expected = Counter({"no2": 720, "so2": 720}) if station == "sapu" else Counter({"no2": 720})
                self.assertEqual(Counter(row["pollutant"] for row in result["rows"]), expected)
                self.assertEqual({row["station_id"] for row in result["rows"]}, {station})
                self.assertEqual({row["unit"] for row in result["rows"]}, {"µg/m3N"})
                self.assertEqual(min(row["date"] for row in result["rows"]), "2025-06-01")
                self.assertEqual(max(row["date"] for row in result["rows"]), "2025-06-30")
                self.assertTrue(any("período declarado" in issue for issue in result["issues"]))
                self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)


if __name__ == "__main__":
    unittest.main()
