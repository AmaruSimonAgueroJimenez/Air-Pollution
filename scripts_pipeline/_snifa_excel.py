"""Read SNIFA hourly gas matrices and timestamped columns, read-only.

Supported inputs are legacy XLS (xlrd) and XLSX/XLSM (openpyxl). A station
must match its configured ``id`` or an ``aliases`` entry in sheet metadata,
the filename, or either of its two immediate parent directories. Receiving
one station configuration alone is not evidence of station identity.

Dates and hour labels remain source-local, including hour 24. No time zone,
quality validation, unit conversion, interpolation or deduplication is applied.
Only CO, SO2, NO2 and O3 hourly matrices are read; eight-hour and particulate
sheets are excluded. Source files and cached formula values are read only.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
import calendar
from contextlib import ExitStack
import math
from numbers import Real
from pathlib import Path
import re
import unicodedata
from typing import Callable


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).casefold()
    return re.sub(r"[^a-z0-9]+", " ", "".join(c for c in text if not unicodedata.combining(c))).strip()


def _raw(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value)


def _column(index: int) -> str:
    result = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


@dataclass
class _Sheet:
    name: str
    rows: list[list]
    serial_date: Callable[[float], datetime]
    reader_note: str = ""


def _read_calamine(path: Path) -> list[_Sheet]:
    """Read XLS rejected by xlrd without changing BIFF strings or source bytes."""
    from python_calamine import CalamineWorkbook

    def explicit_dates_only(value):
        # Calamine supplies formatted date cells as date/datetime objects. Its
        # Python API does not expose the workbook epoch: never guess numeric dates.
        raise ValueError("Fecha serial sin formato no inferida por lector alternativo")

    with CalamineWorkbook.from_path(path) as workbook:
        return [_Sheet(name, workbook.get_sheet_by_name(name).to_python(skip_empty_area=False),
                       explicit_dates_only,
                       "Lectura alternativa python-calamine tras AssertionError de xlrd; "
                       "fechas seriales sin formato no inferidas")
                for name in workbook.sheet_names]


def _read_sheets(path: Path) -> list[_Sheet]:
    with path.open("rb") as stream:
        signature = stream.read(8)
    # SNIFA preserves legacy .XLS names for some actual XLSX/ZIP workbooks.
    # Select the reader by the file signature and keep the source untouched.
    if signature == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        import xlrd

        try:
            workbook = xlrd.open_workbook(str(path), on_demand=True)
        except AssertionError:
            # Some official XLS have SST counts rejected by xlrd. An independent
            # read-only backend handles them; do not patch counts or strings.
            return _read_calamine(path)
        try:
            mode = workbook.datemode
            convert = lambda value: xlrd.xldate_as_datetime(value, mode)
            sheets = []
            for sheet in workbook.sheets():
                rows = []
                for row in sheet.get_rows():
                    rows.append([
                        xlrd.error_text_from_code.get(cell.value, str(cell.value))
                        if cell.ctype == xlrd.XL_CELL_ERROR else
                        bool(cell.value) if cell.ctype == xlrd.XL_CELL_BOOLEAN else
                        cell.value for cell in row
                    ])
                sheets.append(_Sheet(sheet.name, rows, convert))
            return sheets
        finally:
            workbook.release_resources()
    if not signature.startswith(b"PK"):
        raise ValueError("Archivo sin firma de libro Excel OLE/XLS o ZIP/XLSX")
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel

    with ExitStack() as stack:
        # File objects avoid openpyxl's extension check on mislabeled originals.
        cached_stream = stack.enter_context(path.open("rb"))
        workbook = load_workbook(cached_stream, read_only=True, data_only=True)
        stack.callback(workbook.close)
        formula_stream = stack.enter_context(path.open("rb"))
        formulas = load_workbook(formula_stream, read_only=True, data_only=False)
        stack.callback(formulas.close)
        epoch = workbook.epoch
        convert = lambda value: from_excel(value, epoch)
        sheets = []
        for sheet in workbook:
            rows = []
            for cached, source in zip(sheet.iter_rows(values_only=True),
                                      formulas[sheet.title].iter_rows(values_only=True)):
                # An uncached formula is unavailable, never a genuine empty value.
                rows.append([original if value is None and isinstance(original, str)
                             and original.startswith("=") else value
                             for value, original in zip(cached, source)])
            sheets.append(_Sheet(sheet.title, rows, convert))
        return sheets


def _station_matches(text: str, stations: list[dict]) -> set[str]:
    normalized = " " + _norm(text) + " "
    matches = set()
    for station in stations:
        aliases = station.get("aliases", [])
        if isinstance(aliases, str):
            aliases = [aliases]
        for alias in [station.get("id", ""), *aliases]:
            token = _norm(alias)
            if token and " " + token + " " in normalized:
                matches.add(str(station["id"]))
    return matches


def _label_values(metadata: list[list], label: str) -> list[str]:
    values = []
    for row in metadata:
        for column, cell in enumerate(row):
            if not isinstance(cell, str):
                continue
            before, separator, after = cell.partition(":")
            if _norm(before) != label:
                continue
            value = after.strip() if separator else ""
            if not value:
                value = next((_raw(item).strip() for item in row[column + 1:column + 7]
                              if _raw(item).strip() not in {"", ":"}), "")
            if value:
                values.append(value)
    return values


def _station(path: Path, metadata: str, stations: list[dict], metadata_rows: list[list]) -> str | None:
    local = _station_matches(metadata, stations)
    explicit = [_station_matches(value, stations) for value in _label_values(metadata_rows, "estacion")]
    # An explicit unknown station cannot be overridden by a familiar filename.
    if any(len(matches) != 1 for matches in explicit):
        return None
    file_matches = _station_matches(path.name, stations)
    # Use the nearest matching directory, so an outer project directory cannot
    # override the station-specific directory that owns the downloaded file.
    directory = set()
    for parent in list(path.parents)[:2]:
        directory = _station_matches(parent.name, stations)
        if directory:
            break
    candidates = [matches for matches in (local, file_matches, directory, *explicit) if matches]
    if not candidates:
        return None
    common = set.intersection(*candidates)
    return next(iter(common)) if len(common) == 1 else None


_POLLUTANTS = {
    "co": (r"\bco\b", r"\bmonoxido de carbono\b"),
    "so2": (r"\bso2\b", r"\bdioxido de azufre\b"),
    "no2": (r"\bno2\b", r"\bdioxido de nitrogeno\b"),
    "o3": (r"\bo3\b", r"\bozono\b"),
}


def _pollutants(text: str) -> set[str]:
    normalized = _norm(text)
    return {name for name, patterns in _POLLUTANTS.items()
            if any(re.search(pattern, normalized) for pattern in patterns)}


def _excluded(text: str) -> bool:
    normalized = _norm(text)
    return bool(re.search(r"\b(?:8|ocho)\s*(?:h|hr|hrs|horas?)\b|\boctohor\w*|\bmovil\w*|\bmoving\b", normalized))


def _hour(value: object) -> tuple[int, str] | None:
    if isinstance(value, time):
        return (value.hour, value.strftime("%H%M")) if not (value.minute or value.second or value.microsecond) else None
    raw = _raw(value).strip()
    if isinstance(value, bool):
        return None
    if isinstance(value, Real) and 0 < value < 1:
        hour = float(value) * 24
        return (round(hour), raw) if abs(hour - round(hour)) < 1e-7 else None
    if re.fullmatch(r"\d{1,2}:00(?::00)?", raw):
        number = int(raw.split(":")[0])
    elif re.fullmatch(r"\d{1,4}", raw):
        number = int(raw)
        if number > 24 or len(raw) >= 3:
            if number % 100:
                return None
            number //= 100
    else:
        return None
    return (number, raw) if 0 <= number <= 24 else None


def _headers(rows: list[list], convert: Callable) -> list[tuple[int, list[tuple[int, str]]]]:
    headers = []
    sequences = [list(range(1, 25)), list(range(24)), list(range(25))]
    for row_index, row in enumerate(rows):
        for start in range(max(0, len(row) - 23)):
            # Real observations can coincidentally equal 0..23 or 1..24.
            # A date to the left identifies a data row, never an hour header.
            if any(_date(value, convert) is not None for value in row[:start]):
                continue
            for sequence in reversed(sequences):
                if start + len(sequence) > len(row):
                    continue
                parsed = [_hour(value) for value in row[start:start + len(sequence)]]
                if all(item is not None for item in parsed) and [item[0] for item in parsed] == sequence:
                    headers.append((row_index, [(start + i, item[1]) for i, item in enumerate(parsed)]))
                    break
            else:
                continue
            break
    return headers


def _date(value: object, convert: Callable) -> date | None:
    try:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = _raw(value).strip()
        if re.fullmatch(r"\d{8}", text):
            return datetime.strptime(text, "%Y%m%d").date()
        if isinstance(value, Real) and not isinstance(value, bool) and 1 <= value <= 100000:
            converted = convert(float(value))
            return converted.date() if isinstance(converted, datetime) else None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
        return datetime.fromisoformat(text).date() if re.match(r"\d{4}-\d{2}-\d{2}[ T]", text) else None
    except (ValueError, OverflowError, TypeError):
        return None


def _date_column(rows: list[list], header: int, first_hour: int, convert: Callable) -> int | None:
    counts = {column: sum(_date(row[column], convert) is not None for row in rows[header + 1:header + 36]
                          if len(row) > column) for column in range(first_hour)}
    best = max(counts.values(), default=0)
    winners = [column for column, count in counts.items() if count == best and count]
    return winners[0] if len(winners) == 1 else None


_UNIT = re.compile(r"(?<!\w)(?:ppmv?|ppbv?|(?:[µμu]g|mg)\s*/\s*m(?:3|³|\^3)\s*[Nn]?)(?!\w)", re.I)


def _unit(metadata: list[list]) -> str:
    explicit = []
    for row in metadata:
        for column, value in enumerate(row):
            if not isinstance(value, str):
                continue
            match = re.match(r"\s*(?:UNIDAD(?:ES)?|UNIT(?:S)?)(?:\s*:\s*(.*)|\s*)$", value, re.I)
            if not match:
                continue
            remainder = (match.group(1) or "").strip()
            if not remainder:
                remainder = next((_raw(cell).strip() for cell in row[column + 1:column + 7]
                                  if _raw(cell).strip() not in {"", ":"}), "")
            if remainder:
                explicit.append(remainder)
    units = set(explicit)
    if not units:
        units = {match.group(0).strip() for row in metadata for value in row
                 for match in _UNIT.finditer(_raw(value))}
    return next(iter(units)) if len(units) == 1 else "unknown"


_MONTHS = {month: number for number, month in enumerate(
    ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"), 1)}


def _declared_period(metadata: str) -> tuple[date, date] | None:
    match = re.search(r"\b(\d{1,2})\s+(?:al|a)\s+(\d{1,2})\s+de\s+(\w+)\s+(?:del|de)\s+(\d{4})\b", _norm(metadata))
    if match and match[3] in _MONTHS:
        try:
            return date(int(match[4]), _MONTHS[match[3]], int(match[1])), date(int(match[4]), _MONTHS[match[3]], int(match[2]))
        except ValueError:
            pass
    month = re.search(r"\bperiodo\s+(\w+)\s+(?:del|de)\s+(\d{4})\b", _norm(metadata))
    if month and month[1] in _MONTHS:
        year, number = int(month[2]), _MONTHS[month[1]]
        try:
            return date(year, number, 1), date(year, number, calendar.monthrange(year, number)[1])
        except ValueError:
            pass
    return None


def _value(value: object) -> tuple[float | None, str, str]:
    raw = _raw(value)
    if raw.strip() == "":
        return None, raw, ""
    if isinstance(value, Real) and not isinstance(value, bool) and math.isfinite(value):
        return float(value), raw, ""
    numeric = raw.strip().replace("−", "-")
    if re.fullmatch(r"[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:[eE][+-]?\d+)?", numeric):
        parsed = float(numeric.replace(",", "."))
        if math.isfinite(parsed):
            return parsed, raw, ""
    return None, raw, "FORMULA_NO_CACHED_VALUE" if raw.startswith("=") else raw.strip()


def _timestamp(value: object, convert: Callable) -> tuple[date, str] | None:
    """Keep the published day/hour, including an explicit 2400, without UTC conversion."""
    try:
        raw = _raw(value).strip()
        compact = re.fullmatch(r"(\d{8})\s*(\d{4})", raw)
        if compact:
            day = datetime.strptime(compact[1], "%Y%m%d").date()
            hour = _hour(compact[2])
            return (day, compact[2]) if hour is not None else None
        if isinstance(value, datetime):
            stamp = value
        elif isinstance(value, Real) and not isinstance(value, bool) and 1 <= value <= 100000:
            stamp = convert(float(value))
        else:
            stamp = None
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S",
                        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"):
                try:
                    stamp = datetime.strptime(raw, fmt)
                    break
                except ValueError:
                    pass
        if isinstance(stamp, datetime) and stamp.tzinfo is None and not (
                stamp.minute or stamp.second or stamp.microsecond):
            return stamp.date(), stamp.strftime("%H%M")
    except (ValueError, OverflowError, TypeError):
        pass
    return None


def _columnar_layout(sheet: _Sheet, path: Path, stations: list[dict], events=False):
    """Recognize the explicit two-row 'Fecha y hora' layout used by Los Guindos."""
    for index, row in enumerate(sheet.rows[:40]):
        date_columns = [i for i, value in enumerate(row)
                        if re.fullmatch(r"fecha(?: y)? hora", _norm(value))]
        if len(date_columns) != 1 or index + 1 >= len(sheet.rows):
            continue
        labels = sheet.rows[index + 1]
        columns = []
        problems = []
        for column, label in enumerate(labels):
            pollutants = _pollutants(_raw(label))
            kind = _raw(row[column]) if column < len(row) else ""
            description = _raw(label) + " " + kind
            if not pollutants or _excluded(description):
                continue
            if len(pollutants) != 1 or re.search(r"\b(?:no|nox)\b", _norm(label)):
                problems.append(f"{sheet.name}: columna {_column(column)} de contaminante ambiguo omitida")
                continue
            if events:
                if "invalidacion" not in _norm(kind):
                    continue
            elif "horari" not in _norm(kind):
                problems.append(f"{sheet.name}: columna {_column(column)} sin frecuencia horaria explícita omitida")
                continue
            unit = _unit([[label]]) if not events else ""
            columns.append((column, next(iter(pollutants)), unit))
        if not columns:
            return None, problems
        repeated = {pollutant for _, pollutant, _ in columns
                    if sum(other == pollutant for _, other, _ in columns) > 1}
        if repeated:
            problems.append(f"{sheet.name}: columnas horarias duplicadas/ambiguas {', '.join(sorted(repeated))} omitidas")
            columns = [column for column in columns if column[1] not in repeated]
        metadata_rows = list(sheet.rows[:index])
        # These workbooks publish station identity as a sentence, not a label cell.
        for metadata_row in sheet.rows[:index]:
            for cell in metadata_row:
                match = re.fullmatch(r"monitoreos\s+estacion\s+(.+)", _norm(cell))
                if match:
                    metadata_rows.append(["estacion", match.group(1)])
        metadata = " ".join([sheet.name, *(_raw(v) for line in metadata_rows for v in line)])
        # A column table with several gases must carry its own station identity.
        station = _station(path, metadata, stations, metadata_rows) \
            if len(_station_matches(metadata, stations)) == 1 else None
        if station is None:
            return None, [*problems, f"{sheet.name}: estación no identificada inequívocamente; columnas omitidas"]
        if not events:
            for column, pollutant, unit in columns:
                if unit == "unknown":
                    problems.append(f"{sheet.name}: unidad ausente o ambigua para {pollutant}; se conserva como unknown")
        return {"header": index, "date_column": date_columns[0], "columns": columns,
                "station": station, "declared": _declared_period(metadata)}, problems
    return None, []


def _parse_columnar(sheet: _Sheet, path: Path, stations: list[dict], sheets: list[_Sheet],
                    start: date | None, end: date | None) -> tuple[list[dict], list[str], bool]:
    layout, issues = _columnar_layout(sheet, path, stations)
    if layout is None:
        return [], issues, bool(issues)
    event_codes = {}
    for event_sheet in sheets:
        if not re.search(r"\beventos\b", _norm(event_sheet.name)):
            continue
        event_layout, _ = _columnar_layout(event_sheet, path, stations, events=True)
        if event_layout is None or event_layout["station"] != layout["station"]:
            continue
        for index, source in enumerate(event_sheet.rows[event_layout["header"] + 2:],
                                       event_layout["header"] + 2):
            column = event_layout["date_column"]
            stamp = _timestamp(source[column], event_sheet.serial_date) if len(source) > column else None
            if stamp is None:
                continue
            for column, pollutant, _ in event_layout["columns"]:
                code = _raw(source[column]).strip() if len(source) > column else ""
                if code:
                    event_codes.setdefault((*stamp, pollutant), []).append(
                        (code, f"{event_sheet.name}!{_column(column)}{index + 1}"))
    rows, outside, invalid_time = [], 0, 0
    for index, source in enumerate(sheet.rows[layout["header"] + 2:], layout["header"] + 2):
        column = layout["date_column"]
        published = source[column] if len(source) > column else None
        stamp = _timestamp(published, sheet.serial_date)
        if stamp is None:
            if _raw(published).strip():
                invalid_time += 1
            continue
        day, hour = stamp
        declared = layout["declared"]
        if declared and not declared[0] <= day <= declared[1]:
            outside += 1
            continue
        if (start and day < start) or (end and day > end):
            continue
        for column, pollutant, unit in layout["columns"]:
            value, raw, code = _value(source[column] if len(source) > column else None)
            locator = f"{sheet.name}!{_column(column)}{index + 1}"
            flags = event_codes.get((day, hour, pollutant), [])
            codes = [code] if code else []
            for flag, event_locator in flags:
                if flag not in codes:
                    codes.append(flag)
                locator += "; " + event_locator
            rows.append({"station_id": layout["station"], "pollutant": pollutant,
                         "date": day.isoformat(), "hour_label": hour, "value": value,
                         "raw_value": raw, "unit": unit, "quality_code": " | ".join(codes),
                         "source_locator": locator, "resolution": "hourly"})
    if outside:
        issues.append(f"{sheet.name}: {outside} filas de fecha fuera del período declarado en el archivo fueron excluidas")
    if invalid_time:
        issues.append(f"{sheet.name}: {invalid_time} filas sin sello horario entero reconocible fueron excluidas")
    return rows, issues, True


def _parse_validated_hourly(sheet: _Sheet, path: Path, stations: list[dict], sheets: list[_Sheet],
                            start: date | None, end: date | None):
    """Read explicit hourly validated value/code columns, never five-minute raw data."""
    if not re.search(r"\bdata horaria validada\b", _norm(sheet.name)):
        return [], [], False
    pollutants = _pollutants(sheet.name)
    if len(pollutants) != 1:
        return [], [], True
    header = next((i for i, row in enumerate(sheet.rows[:20]) if len(row) >= 3
                   and _norm(row[0]) == "aaaammdd hhmm"
                   and "codigo de invalidacion" in _norm(row[2])), None)
    if header is None:
        return [], [f"{sheet.name}: columnas horarias validadas sin cabecera inequívoca; omitidas"], True
    metadata_rows = list(sheet.rows[:header + 1])
    # El Peñón publishes station identity once on its summary sheet. Preserve
    # explicit conflicting/unknown declarations rather than trusting a filename.
    for other in sheets:
        for row in other.rows[:30]:
            if row and _norm(row[0]) in {"nombre de la estacion", "estacion"}:
                metadata_rows.append(["estacion", *row[1:]])
    metadata = " ".join(_raw(value) for row in metadata_rows for value in row)
    station = _station(path, metadata, stations, metadata_rows)
    if station is None:
        return [], [f"{sheet.name}: estación no identificada inequívocamente; columnas omitidas"], True
    unit = _unit([[sheet.rows[header][1]]])
    issues = [] if unit != "unknown" else [
        f"{sheet.name}: unidad ausente o ambigua; se conserva como unknown"]
    rows, invalid_time = [], 0
    for index, source in enumerate(sheet.rows[header + 1:], header + 1):
        if not source:
            continue
        stamp = _timestamp(source[0], sheet.serial_date)
        if stamp is None:
            if _raw(source[0]).strip():
                invalid_time += 1
            continue
        day, hour = stamp
        if (start and day < start) or (end and day > end):
            continue
        value, raw, code = _value(source[1] if len(source) > 1 else None)
        flag = _raw(source[2]).strip() if len(source) > 2 else ""
        codes = list(dict.fromkeys(item for item in (code, flag) if item))
        locator = f"{sheet.name}!B{index + 1}"
        if flag:
            locator += f"; {sheet.name}!C{index + 1}"
        rows.append({"station_id": station, "pollutant": next(iter(pollutants)),
                     "date": day.isoformat(), "hour_label": hour, "value": value,
                     "raw_value": raw, "unit": unit, "quality_code": " | ".join(codes),
                     "source_locator": locator, "resolution": "hourly"})
    if invalid_time:
        issues.append(f"{sheet.name}: {invalid_time} filas sin sello horario entero reconocible fueron excluidas")
    return rows, issues, True


def parse_excel(path: Path, stations: list[dict], period_start: str = "", period_end: str = "") -> dict:
    """Return observation rows plus actionable issues; never write the input.

    Optional date bounds are inclusive and intersect a confidently recognized
    source-declared period. Non-numeric source codes and empty hourly cells are
    retained as null values. Units are original labels, or ``unknown`` with an
    issue. Unsupported/ambiguous tables are reported instead of guessed.
    """
    path = Path(path)
    rows, issues = [], []
    start = date.fromisoformat(period_start) if period_start else None
    end = date.fromisoformat(period_end) if period_end else None
    if start and end and start > end:
        raise ValueError("period_start debe ser anterior o igual a period_end")
    try:
        sheets = _read_sheets(path)
    except Exception as exc:
        # Reader backends use different exception families for malformed files.
        # Report failure for this source so a batch can keep its other sources.
        return {"rows": [], "issues": [f"No se pudo leer {path.name}: {type(exc).__name__}: {exc}"]}
    issues.extend(sheet.reader_note for sheet in sheets if sheet.reader_note)
    for sheet in sheets:
        if re.search(r"\beventos\b|\bdata cruda\b", _norm(sheet.name)):
            continue
        if _excluded(sheet.name) or re.search(r"\b(?:mp\s*10|pm\s*10|mp\s*2\s*5|pm\s*2\s*5|no|nox)\b", _norm(sheet.name)):
            continue
        validated_rows, validated_issues, validated_layout = _parse_validated_hourly(
            sheet, path, stations, sheets, start, end)
        rows.extend(validated_rows)
        issues.extend(validated_issues)
        if validated_layout:
            continue
        column_rows, column_issues, column_layout = _parse_columnar(sheet, path, stations, sheets, start, end)
        rows.extend(column_rows)
        issues.extend(column_issues)
        if column_layout:
            continue
        headers = _headers(sheet.rows, sheet.serial_date)
        if not headers:
            if _pollutants(sheet.name):
                issues.append(f"{sheet.name}: no se encontró una matriz horaria reconocible")
            continue
        for index, (header, hours) in enumerate(headers):
            previous = headers[index - 1][0] + 1 if index else 0
            metadata_rows = sheet.rows[previous:header]
            metadata = " ".join([sheet.name, *(_raw(value) for row in metadata_rows for value in row)])
            if _excluded(metadata):
                continue
            variables = _label_values(metadata_rows, "variable")
            if variables and any(len(_pollutants(value)) != 1 for value in variables):
                continue
            pollutants = _pollutants(metadata)
            if len(pollutants) != 1:
                if pollutants:
                    issues.append(f"{sheet.name}: contaminante ambiguo; matriz omitida")
                continue
            station = _station(path, metadata, stations, metadata_rows)
            if station is None:
                issues.append(f"{sheet.name}: estación no identificada inequívocamente; matriz omitida")
                continue
            column = _date_column(sheet.rows, header, hours[0][0], sheet.serial_date)
            if column is None:
                issues.append(f"{sheet.name}: columna de fechas ausente o ambigua; matriz omitida")
                continue
            unit = _unit(metadata_rows)
            if unit == "unknown":
                issues.append(f"{sheet.name}: unidad ausente o ambigua; se conserva como unknown")
            declared = _declared_period(metadata)
            excluded = 0
            stop = headers[index + 1][0] if index + 1 < len(headers) else len(sheet.rows)
            for row_index in range(header + 1, stop):
                source = sheet.rows[row_index]
                day = _date(source[column], sheet.serial_date) if len(source) > column else None
                if day is None:
                    continue
                if declared and not declared[0] <= day <= declared[1]:
                    excluded += 1
                    continue
                if (start and day < start) or (end and day > end):
                    continue
                for hour_column, hour in hours:
                    value, raw, code = _value(source[hour_column] if len(source) > hour_column else None)
                    rows.append({"station_id": station, "pollutant": next(iter(pollutants)),
                                 "date": day.isoformat(), "hour_label": hour, "value": value,
                                 "raw_value": raw, "unit": unit, "quality_code": code,
                                 "source_locator": f"{sheet.name}!{_column(hour_column)}{row_index + 1}",
                                 "resolution": "hourly"})
            if excluded:
                issues.append(f"{sheet.name}: {excluded} filas de fecha fuera del período declarado en el archivo fueron excluidas")
    if not rows and not issues:
        issues.append(f"{path.name}: no se encontraron matrices horarias de gases dentro del alcance")
    return {"rows": rows, "issues": list(dict.fromkeys(issues))}
